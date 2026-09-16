"""Формирование повестки заседания в DOCX.

Если загружен шаблон DocumentTemplate вида «agenda» (для комиссии или общий) —
используется docxtpl. Контекст шаблона:
  organization, commission (name, short_name), meeting_date, meeting_time, weekday,
  place, format, video_link, kind, agenda_status, approved_by, approved_at,
  chairman, secretary,
  items — список {number, title, description, speaker, duration, patient_id, rubric, requires_vote}.
Иначе формируется документ по умолчанию (python-docx).
"""
import io

from django.conf import settings
from django.utils import timezone

from apps.committees.models import DocumentTemplate

WEEKDAYS = ["понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье"]
MONTHS_GEN = [
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
]


def date_long(d):
    return f"{d.day} {MONTHS_GEN[d.month - 1]} {d.year} г."


def agenda_context(meeting):
    local = timezone.localtime(meeting.starts_at)
    commission = meeting.commission
    show_patient = commission.handles_patient_cases
    items = []
    for item in meeting.agenda_items.select_related("speaker", "rubric").order_by("position", "id"):
        items.append({
            "number": item.position,
            "title": item.title,
            "description": item.description,
            "speaker": item.speaker_display or "",
            "duration": item.duration_minutes,
            "patient_id": item.patient_id if show_patient else "",
            "rubric": item.rubric.name if item.rubric_id else "",
            "requires_vote": item.requires_vote,
            "is_control_item": item.is_control_item,
        })
    chairman = commission.officer("chairman", local.date())
    secretary = commission.officer("secretary", local.date())
    return {
        "organization": settings.ORGANIZATION_NAME,
        "commission": {"name": commission.name, "short_name": commission.short_name},
        "meeting_date": local.strftime("%d.%m.%Y"),
        "meeting_date_long": date_long(local.date()),
        "meeting_time": local.strftime("%H:%M"),
        "weekday": WEEKDAYS[local.weekday()],
        "place": meeting.place,
        "format": meeting.get_format_display(),
        "video_link": meeting.video_link,
        "kind": meeting.get_kind_display(),
        "agenda_status": meeting.get_agenda_status_display(),
        "approved_by": meeting.agenda_approved_by.short_name if meeting.agenda_approved_by_id else "",
        "approved_at": timezone.localtime(meeting.agenda_approved_at).strftime("%d.%m.%Y") if meeting.agenda_approved_at else "",
        "chairman": chairman.short_name if chairman else "",
        "secretary": secretary.short_name if secretary else "",
        "items": items,
        "show_patient_id": show_patient,
    }


def _from_template(template, context):
    from docxtpl import DocxTemplate

    with template.file.open("rb") as fh:
        doc = DocxTemplate(io.BytesIO(fh.read()))
    doc.render(context)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _default(context):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    def para(text, bold=False, align=None, size=None, space_after=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        return p

    if context["approved_by"]:
        p = para("УТВЕРЖДАЮ", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
        para(f"Председатель комиссии {context['approved_by']}", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
        para(context["approved_at"], align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(context["organization"], align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(context["commission"]["name"], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("ПОВЕСТКА ЗАСЕДАНИЯ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    if context["kind"] and context["kind"] != "Плановое":
        para(f"({context['kind'].lower()} заседание)", align=WD_ALIGN_PARAGRAPH.CENTER)

    table = doc.add_table(rows=0, cols=2)
    rows = [
        ("Дата", f"{context['meeting_date_long']} ({context['weekday']})"),
        ("Время", context["meeting_time"]),
        ("Формат", context["format"]),
    ]
    if context["place"]:
        rows.append(("Место", context["place"]))
    if context["video_link"]:
        rows.append(("Видеоконференция", context["video_link"]))
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    doc.add_paragraph()

    if not context["items"]:
        para("Вопросы повестки не сформированы.")
    for item in context["items"]:
        p = doc.add_paragraph()
        run = p.add_run(f"{item['number']}. {item['title']}")
        run.bold = True
        p.paragraph_format.space_after = Pt(2)
        details = []
        if item["patient_id"]:
            details.append(f"ID пациента (МИС): {item['patient_id']}")
        if item["speaker"]:
            details.append(f"Докладчик: {item['speaker']}")
        details.append(f"Регламент: {item['duration']} мин")
        if item["requires_vote"]:
            details.append("выносится на голосование")
        para("; ".join(details), size=11, space_after=2)
        if item["description"]:
            para(item["description"], size=11)

    doc.add_paragraph()
    if context["secretary"]:
        para(f"Секретарь комиссии ____________ {context['secretary']}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def agenda_docx(meeting):
    """Возвращает (bytes, имя файла)."""
    context = agenda_context(meeting)
    template = DocumentTemplate.find(DocumentTemplate.Kind.AGENDA, meeting.commission)
    data = _from_template(template, context) if template else _default(context)
    filename = f"Повестка_{meeting.commission.short_name}_{timezone.localtime(meeting.starts_at):%Y-%m-%d}.docx"
    return data, filename
