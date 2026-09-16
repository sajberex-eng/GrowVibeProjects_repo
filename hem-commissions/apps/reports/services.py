"""Отчёт о работе комиссий: расчёт показателей, выгрузка в XLSX и PDF, подписание."""
import io
from collections import OrderedDict, defaultdict
from datetime import date, datetime, time, timedelta
from xml.sax.saxutils import escape

from django.conf import settings
from django.db.models import Q
from django.utils import timezone

from apps.committees.models import CompositionChange, LegalAct, MemberRecord, Order
from apps.meetings.models import AgendaItem, Attendance, Meeting
from apps.protocols.models import Assignment, Decision, DissentingOpinion, Protocol
from apps.voting.models import VotingSession

NO_RUBRIC = "Без рубрики"


# ---------- периоды ----------

PERIOD_CHOICES = [
    ("h1", "1-е полугодие"),
    ("h2", "2-е полугодие"),
    ("year", "Календарный год"),
    ("custom", "Произвольный период"),
]


def period_bounds(preset, year):
    if preset == "h1":
        return date(year, 1, 1), date(year, 6, 30)
    if preset == "h2":
        return date(year, 7, 1), date(year, 12, 31)
    return date(year, 1, 1), date(year, 12, 31)


def default_period(today=None):
    today = today or timezone.localdate()
    return ("h1" if today.month <= 6 else "h2"), today.year


def _range(date_from, date_to):
    start = timezone.make_aware(datetime.combine(date_from, time.min))
    end = timezone.make_aware(datetime.combine(date_to + timedelta(days=1), time.min))
    return start, end


def _pct(part, whole):
    return round(part * 100 / whole, 1) if whole else None


# ---------- расчёт ----------

def _meeting_section(commission, start, end):
    meetings = list(
        Meeting.objects.filter(commission=commission, starts_at__gte=start, starts_at__lt=end)
        .prefetch_related("reschedules")
        .order_by("starts_at")
    )
    rows = []
    counts = {
        "total": 0, "held": 0, "postponed": 0, "cancelled": 0, "upcoming": 0,
        "extra": 0, "extra_held": 0, "quorum_reached": 0, "quorum_not_reached": 0,
    }
    for m in meetings:
        rescheduled = bool(m.reschedules.all()) or m.status == Meeting.Status.POSTPONED
        quorum = m.quorum() if m.status == Meeting.Status.HELD else None
        counts["total"] += 1
        if m.status == Meeting.Status.HELD:
            counts["held"] += 1
            counts["quorum_reached" if quorum["reached"] else "quorum_not_reached"] += 1
        elif m.status == Meeting.Status.CANCELLED:
            counts["cancelled"] += 1
        else:
            counts["upcoming"] += 1
        if rescheduled:
            counts["postponed"] += 1
        if m.kind == Meeting.Kind.EXTRA:
            counts["extra"] += 1
            if m.status == Meeting.Status.HELD:
                counts["extra_held"] += 1
        rows.append({"meeting": m, "rescheduled": rescheduled, "quorum": quorum})
    return meetings, rows, counts


def _attendance_section(commission, held):
    stats = OrderedDict()
    role_names = dict(MemberRecord.Role.choices)
    for m in held:
        records = list(m.composition())
        marks = {a.user_id: a.status for a in Attendance.objects.filter(meeting=m)}
        for rec in records:
            row = stats.get(rec.user_id)
            if row is None:
                row = stats[rec.user_id] = {
                    "user": rec.user, "name": rec.user.full_name or rec.user.username, "role": role_names.get(rec.role, ""),
                    "meetings": 0, "present": 0, "online": 0, "absent_excused": 0, "absent": 0, "unmarked": 0,
                }
            row["role"] = role_names.get(rec.role, row["role"])
            row["meetings"] += 1
            status = marks.get(rec.user_id)
            if status == Attendance.Status.PRESENT:
                row["present"] += 1
            elif status == Attendance.Status.ONLINE:
                row["online"] += 1
            elif status == Attendance.Status.ABSENT_EXCUSED:
                row["absent_excused"] += 1
            elif status == Attendance.Status.ABSENT:
                row["absent"] += 1
            else:
                row["unmarked"] += 1
    rows = list(stats.values())
    for row in rows:
        row["percent"] = _pct(row["present"] + row["online"], row["meetings"])
    rows.sort(key=lambda r: r["name"])
    attended = sum(r["present"] + r["online"] for r in rows)
    slots = sum(r["meetings"] for r in rows)
    return rows, _pct(attended, slots)


def _protocol_section(commission, held, start_date, end_date):
    days = []
    unsigned = 0
    protocols = (
        Protocol.objects.filter(commission=commission, kind=Protocol.Kind.MEETING, meeting__in=held)
        .exclude(status=Protocol.Status.ANNULLED)
        .select_related("meeting")
    )
    for p in protocols:
        if p.status == Protocol.Status.SIGNED and p.signed_at:
            days.append((timezone.localtime(p.signed_at).date() - p.meeting.meeting_date).days)
        else:
            unsigned += 1
    timeliness = {
        "signed": len(days),
        "unsigned": unsigned,
        "avg_days": round(sum(days) / len(days), 1) if days else None,
        "max_days": max(days) if days else None,
    }
    dissents = DissentingOpinion.objects.filter(
        protocol__commission=commission, protocol__protocol_date__gte=start_date,
        protocol__protocol_date__lte=end_date, withdrawn_at__isnull=True,
    ).exclude(protocol__status=Protocol.Status.ANNULLED).count()
    return timeliness, dissents


def _rubric_section(held):
    counter = defaultdict(int)
    items = AgendaItem.objects.filter(meeting__in=held, is_control_item=False).select_related("rubric")
    total = 0
    for item in items:
        counter[item.rubric.name if item.rubric_id else NO_RUBRIC] += 1
        total += 1
    rows = sorted(counter.items(), key=lambda kv: (kv[0] == NO_RUBRIC, -kv[1], kv[0]))
    return [{"name": k, "count": v} for k, v in rows], total


def _assignment_state(a, today):
    if a.status == Assignment.Status.CANCELLED:
        return "cancelled"
    if a.status == Assignment.Status.DONE_LATE:
        return "done_late"
    if a.status == Assignment.Status.DONE:
        if a.completed_on and a.completed_on > a.due_date:
            return "done_late"
        return "done_on_time"
    if a.due_date < today:
        return "overdue"
    return "open"


ASSIGNMENT_STATES = OrderedDict([
    ("done_on_time", "Исполнено в срок"),
    ("done_late", "Исполнено с нарушением срока"),
    ("overdue", "Просрочено"),
    ("open", "В работе (срок не истёк)"),
    ("cancelled", "Снято"),
])


def _assignment_section(commission, start_date, end_date, today):
    qs = (
        Assignment.objects.filter(commission=commission)
        .filter(
            Q(decision__protocol__protocol_date__gte=start_date, decision__protocol__protocol_date__lte=end_date)
            | Q(due_date__gte=start_date, due_date__lte=end_date)
        )
        .exclude(decision__protocol__status=Protocol.Status.ANNULLED)
        .select_related("responsible", "decision", "decision__protocol")
        .distinct()
        .order_by("due_date", "id")
    )
    totals = {key: 0 for key in ASSIGNMENT_STATES}
    totals["total"] = 0
    by_resp = OrderedDict()
    rows = []
    for a in qs:
        state = _assignment_state(a, today)
        totals[state] += 1
        totals["total"] += 1
        name = a.responsible_display or "Не указан"
        resp = by_resp.setdefault(name, {"name": name, "total": 0, **{k: 0 for k in ASSIGNMENT_STATES}})
        resp[state] += 1
        resp["total"] += 1
        rows.append({"assignment": a, "state": state, "state_label": ASSIGNMENT_STATES[state], "responsible": name})
    return totals, sorted(by_resp.values(), key=lambda r: r["name"]), rows


def _legal_section(commission):
    acts = list(LegalAct.objects.filter(commission=commission).order_by("sort", "id"))
    counts = {key: 0 for key, _ in LegalAct.Status.choices}
    for act in acts:
        counts[act.status] += 1
    checks = [a.last_checked_at for a in acts if a.last_checked_at]
    return {
        "total": len(acts),
        "actual": counts[LegalAct.Status.ACTUAL],
        "lost_force": counts[LegalAct.Status.LOST_FORCE] + counts[LegalAct.Status.RETIRED],
        "needs_check": counts[LegalAct.Status.NEEDS_CHECK] + counts[LegalAct.Status.UNAVAILABLE],
        "last_check": max(checks) if checks else None,
        "acts": acts,
    }


def build_commission_report(commission, date_from, date_to, today=None):
    today = today or timezone.localdate()
    start, end = _range(date_from, date_to)
    meetings, meeting_rows, meeting_counts = _meeting_section(commission, start, end)
    held = [m for m in meetings if m.status == Meeting.Status.HELD]
    attendance, avg_attendance = _attendance_section(commission, held)
    timeliness, dissents = _protocol_section(commission, held, date_from, date_to)
    rubrics, questions_total = _rubric_section(held)
    decisions = list(
        Decision.objects.filter(
            protocol__commission=commission, protocol__protocol_date__gte=date_from, protocol__protocol_date__lte=date_to
        )
        .exclude(protocol__status__in=[Protocol.Status.ANNULLED])
        .select_related("protocol", "item")
        .order_by("protocol__protocol_date", "item__position", "id")
    )
    assignments, by_responsible, assignment_rows = _assignment_section(commission, date_from, date_to, today)
    absentee = VotingSession.objects.filter(
        commission=commission, meeting__isnull=True, status=VotingSession.Status.CLOSED,
        closed_at__gte=start, closed_at__lt=end,
    ).count()
    changes = list(
        CompositionChange.objects.filter(
            commission=commission, status=CompositionChange.Status.APPLIED,
            applied_on__gte=date_from, applied_on__lte=date_to,
        ).select_related("user", "order").order_by("applied_on", "id")
    )
    orders = list(
        Order.objects.filter(
            commission=commission, status=Order.Status.SIGNED, signed_on__gte=date_from, signed_on__lte=date_to
        ).order_by("signed_on")
    )
    return {
        "commission": commission,
        "chairman": commission.officer(MemberRecord.Role.CHAIRMAN, min(date_to, today)),
        "meetings": meeting_counts,
        "meeting_rows": meeting_rows,
        "absentee_votings": absentee,
        "attendance": attendance,
        "avg_attendance": avg_attendance,
        "protocol_timeliness": timeliness,
        "dissents": dissents,
        "rubrics": rubrics,
        "questions_total": questions_total,
        "decisions": decisions,
        "decisions_count": len(decisions),
        "assignments": assignments,
        "assignments_by_responsible": by_responsible,
        "assignment_rows": assignment_rows,
        "legal": _legal_section(commission),
        "composition_changes": changes,
        "orders": orders,
    }


SUM_KEYS = ["total", "held", "postponed", "cancelled", "upcoming", "extra", "extra_held", "quorum_reached", "quorum_not_reached"]


def build_report(commissions, date_from, date_to, today=None):
    """Отчёт по списку комиссий за период [date_from; date_to].

    Возвращает {"date_from", "date_to", "commissions": [...], "totals": {...}}.
    """
    if date_from > date_to:
        date_from, date_to = date_to, date_from
    sections = [build_commission_report(c, date_from, date_to, today=today) for c in commissions]
    totals = {
        "meetings": {k: sum(s["meetings"][k] for s in sections) for k in SUM_KEYS},
        "absentee_votings": sum(s["absentee_votings"] for s in sections),
        "questions_total": sum(s["questions_total"] for s in sections),
        "decisions_count": sum(s["decisions_count"] for s in sections),
        "dissents": sum(s["dissents"] for s in sections),
        "assignments": {
            k: sum(s["assignments"][k] for s in sections) for k in ["total", *ASSIGNMENT_STATES.keys()]
        },
        "legal_total": sum(s["legal"]["total"] for s in sections),
        "composition_changes": sum(len(s["composition_changes"]) for s in sections),
    }
    return {"date_from": date_from, "date_to": date_to, "commissions": sections, "totals": totals}


def report_title(report):
    sections = report["commissions"]
    if len(sections) == 1:
        return f"Отчёт о работе: {sections[0]['commission'].name}"
    return "Отчёт о работе комиссий Центра"


def period_label(report):
    return f"за период с {report['date_from']:%d.%m.%Y} по {report['date_to']:%d.%m.%Y}"


# ---------- XLSX ----------

def _dt(value):
    if not value:
        return ""
    if isinstance(value, datetime):
        return timezone.localtime(value).strftime("%d.%m.%Y %H:%M")
    return value.strftime("%d.%m.%Y")


def export_xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    sheets = OrderedDict()

    def sheet(title, header):
        ws = wb.create_sheet(title)
        ws.append([f"{report_title(report)} {period_label(report)}"])
        ws["A1"].font = Font(bold=True)
        ws.append([])
        ws.append(header)
        for cell in ws[3]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheets[title] = ws
        return ws

    wb.remove(wb.active)
    ws = sheet("Заседания", ["Комиссия", "Дата", "Тип", "Статус", "Переносилось", "Присутствовало", "Состав", "Кворум"])
    for s in report["commissions"]:
        for row in s["meeting_rows"]:
            m, q = row["meeting"], row["quorum"]
            ws.append([
                s["commission"].short_name, _dt(m.starts_at), m.get_kind_display(), m.get_status_display(),
                "да" if row["rescheduled"] else "нет",
                q["present"] if q else "", q["total"] if q else "",
                ("достигнут" if q["reached"] else "не достигнут") if q else "",
            ])
    ws.append([])
    ws.append(["Итого по комиссиям"])
    ws.append(["Комиссия", "Всего", "Проведено", "Переносилось", "Отменено", "Предстоит", "Внеочередных",
               "Внеочередных проведено", "С кворумом", "Без кворума", "Заочных голосований",
               "Ср. посещаемость, %", "Протоколов подписано", "Ср. срок подписания, дн.", "Макс. срок, дн.",
               "Особых мнений"])
    for s in report["commissions"]:
        m, t = s["meetings"], s["protocol_timeliness"]
        ws.append([s["commission"].short_name, m["total"], m["held"], m["postponed"], m["cancelled"], m["upcoming"],
                   m["extra"], m["extra_held"], m["quorum_reached"], m["quorum_not_reached"], s["absentee_votings"],
                   s["avg_attendance"], t["signed"], t["avg_days"], t["max_days"], s["dissents"]])

    ws = sheet("Посещаемость", ["Комиссия", "Член комиссии", "Роль", "Заседаний", "Очно", "Онлайн",
                                "Отсутствовал (уваж.)", "Отсутствовал", "Не отмечено", "Посещаемость, %"])
    for s in report["commissions"]:
        for r in s["attendance"]:
            ws.append([s["commission"].short_name, r["name"], r["role"], r["meetings"], r["present"], r["online"],
                       r["absent_excused"], r["absent"], r["unmarked"], r["percent"]])

    ws = sheet("Вопросы", ["Комиссия", "Рубрика", "Рассмотрено вопросов"])
    for s in report["commissions"]:
        for r in s["rubrics"]:
            ws.append([s["commission"].short_name, r["name"], r["count"]])
        ws.append([s["commission"].short_name, "Итого", s["questions_total"]])

    ws = sheet("Решения", ["Комиссия", "Протокол", "Дата", "№ решения", "Вопрос", "Текст решения"])
    for s in report["commissions"]:
        for d in s["decisions"]:
            ws.append([s["commission"].short_name, d.protocol.number or "б/н", _dt(d.protocol.protocol_date),
                       d.number, d.item.title, d.text])

    ws = sheet("Поручения", ["Комиссия", "Поручение", "Ответственный", "Срок", "Статус", "Дата исполнения", "Оценка"])
    for s in report["commissions"]:
        for r in s["assignment_rows"]:
            a = r["assignment"]
            ws.append([s["commission"].short_name, a.text, r["responsible"], _dt(a.due_date), a.get_status_display(),
                       _dt(a.completed_on), r["state_label"]])
    ws.append([])
    ws.append(["Комиссия", "Ответственный", "Всего", *ASSIGNMENT_STATES.values()])
    for s in report["commissions"]:
        for r in s["assignments_by_responsible"]:
            ws.append([s["commission"].short_name, r["name"], r["total"], *[r[k] for k in ASSIGNMENT_STATES]])

    ws = sheet("НПА", ["Комиссия", "Наименование", "Номер", "Дата", "Статус", "Последняя проверка", "Ссылка"])
    for s in report["commissions"]:
        for act in s["legal"]["acts"]:
            ws.append([s["commission"].short_name, act.title, act.number, _dt(act.act_date), act.get_status_display(),
                       _dt(act.last_checked_at), act.url])

    ws = sheet("Состав", ["Комиссия", "Дата", "Изменение", "Сотрудник", "Роль", "Приказ"])
    for s in report["commissions"]:
        for c in s["composition_changes"]:
            ws.append([s["commission"].short_name, _dt(c.applied_on), c.get_action_display(), str(c.user),
                       c.get_new_role_display() if c.new_role else "", str(c.order) if c.order_id else ""])
    ws.append([])
    ws.append(["Комиссия", "Подписанные приказы", "Дата"])
    for s in report["commissions"]:
        for o in s["orders"]:
            ws.append([s["commission"].short_name, f"{o} «{o.title}»", _dt(o.signed_on)])

    for ws in sheets.values():
        for idx in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(idx)].width = 22
        ws.freeze_panes = "A4"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- PDF ----------

def _p(text, style):
    from reportlab.platypus import Paragraph

    return Paragraph(escape(str(text if text is not None else "—")).replace("\n", "<br/>"), style)


def _table(rows, styles, col_widths=None, header=True):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    data = [[_p(c, styles["small"]) for c in row] for row in rows]
    table = Table(data, colWidths=col_widths, repeatRows=1 if header else 0)
    style = [
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]
    if header:
        style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f4")))
    table.setStyle(TableStyle(style))
    return table


def _fmt(value, suffix=""):
    return "—" if value is None else f"{value}{suffix}"


def export_pdf(report, signature=None):
    """PDF отчёта. signature — {"name", "signed_at", "hash"} для штампа простой ЭП."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, PageBreak, SimpleDocTemplate, Spacer, Table, TableStyle

    from apps.committees.services import pdf_fonts, pdf_styles

    regular, _bold = pdf_fonts()
    st = pdf_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=18 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
        title=report_title(report), author=settings.ORGANIZATION_NAME, invariant=1,
    )
    sections = report["commissions"]
    story = [Spacer(1, 50 * mm), _p(settings.ORGANIZATION_NAME, st["center"]), Spacer(1, 20 * mm)]
    story.append(_p(report_title(report), st["title"]))
    story.append(_p(period_label(report), st["center"]))
    story.append(Spacer(1, 40 * mm))
    if len(sections) == 1:
        chairman = sections[0]["chairman"]
        story.append(_p(
            f"Председатель комиссии ______________________ {chairman.full_name if chairman else ''}", st["normal"]
        ))
    story.append(Spacer(1, 10 * mm))
    story.append(_p(f"Сформирован {timezone.localtime():%d.%m.%Y}", st["small"]))
    story.append(PageBreak())

    width = doc.width
    for s in sections:
        c = s["commission"]
        m, t = s["meetings"], s["protocol_timeliness"]
        story.append(_p(f"{c.name} ({c.short_name})", st["h2"]))
        story.append(_p("1. Формальные показатели", st["bold"]))
        story.append(_table([
            ["Показатель", "Значение"],
            ["Заседаний запланировано в периоде", m["total"]],
            ["Проведено", m["held"]],
            ["Переносилось", m["postponed"]],
            ["Отменено", m["cancelled"]],
            ["Предстоит / не отмечено проведённым", m["upcoming"]],
            ["Из них внеочередных (всего / проведено)", f"{m['extra']} / {m['extra_held']}"],
            ["Заочных голосований", s["absentee_votings"]],
            ["Заседаний с кворумом / без кворума", f"{m['quorum_reached']} / {m['quorum_not_reached']}"],
            ["Средняя посещаемость", _fmt(s["avg_attendance"], " %")],
            ["Протоколов подписано / не подписано", f"{t['signed']} / {t['unsigned']}"],
            ["Срок от заседания до подписания протокола, дн. (средний / макс.)",
             f"{_fmt(t['avg_days'])} / {_fmt(t['max_days'])}"],
            ["Особых мнений", s["dissents"]],
        ], st, [width * 0.7, width * 0.3]))
        story.append(Spacer(1, 4 * mm))
        if s["attendance"]:
            story.append(_p("Посещаемость членов комиссии", st["bold"]))
            rows = [["Член комиссии", "Роль", "Засед.", "Очно", "Онлайн", "Отс. (уваж.)", "Отс.", "%"]]
            for r in s["attendance"]:
                rows.append([r["name"], r["role"], r["meetings"], r["present"], r["online"],
                             r["absent_excused"], r["absent"] + r["unmarked"], _fmt(r["percent"])])
            story.append(_table(rows, st, [width * 0.3, width * 0.18] + [width * 0.0866] * 6))
            story.append(Spacer(1, 4 * mm))

        story.append(_p("2. Содержательные показатели", st["bold"]))
        rows = [["Рубрика", "Рассмотрено вопросов"]] + [[r["name"], r["count"]] for r in s["rubrics"]]
        rows.append(["Итого", s["questions_total"]])
        story.append(_table(rows, st, [width * 0.7, width * 0.3]))
        story.append(Spacer(1, 3 * mm))
        story.append(_p(f"Принято решений: {s['decisions_count']}", st["normal"]))
        if s["decisions"]:
            rows = [["Протокол", "№", "Решение"]]
            for d in s["decisions"]:
                rows.append([f"№ {d.protocol.number or 'б/н'} от {d.protocol.protocol_date:%d.%m.%Y}", d.number, d.text])
            story.append(_table(rows, st, [width * 0.25, width * 0.1, width * 0.65]))
        story.append(Spacer(1, 3 * mm))
        a = s["assignments"]
        story.append(_p("Исполнение поручений", st["bold"]))
        rows = [["Всего", *ASSIGNMENT_STATES.values()], [a["total"], *[a[k] for k in ASSIGNMENT_STATES]]]
        story.append(_table(rows, st, [width / 6] * 6))
        if s["assignments_by_responsible"]:
            rows = [["Ответственный", "Всего", "В срок", "С нарушением", "Просрочено", "В работе", "Снято"]]
            for r in s["assignments_by_responsible"]:
                rows.append([r["name"], r["total"], *[r[k] for k in ASSIGNMENT_STATES]])
            story.append(Spacer(1, 2 * mm))
            story.append(_table(rows, st, [width * 0.34] + [width * 0.11] * 6))
        story.append(Spacer(1, 3 * mm))
        lg = s["legal"]
        story.append(_p("Нормативная правовая база", st["bold"]))
        story.append(_table([
            ["Всего НПА", "Актуальны", "Утратили силу", "Требуют проверки", "Последняя проверка"],
            [lg["total"], lg["actual"], lg["lost_force"], lg["needs_check"], _dt(lg["last_check"]) or "—"],
        ], st, [width / 5] * 5))
        story.append(Spacer(1, 3 * mm))
        story.append(_p("Изменения состава", st["bold"]))
        if s["composition_changes"] or s["orders"]:
            rows = [["Дата", "Изменение", "Сотрудник", "Приказ"]]
            for ch in s["composition_changes"]:
                rows.append([_dt(ch.applied_on), ch.get_action_display(), str(ch.user), str(ch.order) if ch.order_id else ""])
            for o in s["orders"]:
                rows.append([_dt(o.signed_on), "Подписан приказ", o.title, str(o)])
            story.append(_table(rows, st, [width * 0.15, width * 0.25, width * 0.3, width * 0.3]))
        else:
            story.append(_p("Изменений состава в периоде не было.", st["normal"]))
        story.append(Spacer(1, 8 * mm))

    if signature:
        stamp = Table(
            [[_p(
                "ДОКУМЕНТ ПОДПИСАН ПРОСТОЙ ЭЛЕКТРОННОЙ ПОДПИСЬЮ\n"
                f"Подписант: {signature['name']}\n"
                f"Дата и время подписи: {signature['signed_at']}\n"
                f"SHA-256 подписанного документа (без штампа): {signature['hash']}",
                st["small"],
            )]],
            colWidths=[width * 0.8],
        )
        stamp.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#1f4e8c")),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#1f4e8c")),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(KeepTogether([Spacer(1, 6 * mm), stamp]))

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont(regular, 8)
        canvas.drawRightString(A4[0] - 15 * mm, 8 * mm, f"стр. {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ---------- подписание ----------

def signing_purpose(commission, date_from, date_to):
    return f"report:{commission.pk}:{date_from:%Y%m%d}:{date_to:%Y%m%d}"


def sign_report(commission, date_from, date_to, user, ip=None):
    """Формирует PDF отчёта со штампом простой ЭП и сохраняет SignedReport.

    Проверка кода подтверждения выполняется вызывающим кодом.
    """
    from django.core.files.base import ContentFile

    from apps.audit import services as audit
    from apps.committees.services import sha256_bytes

    from .models import SignedReport

    report = build_report([commission], date_from, date_to)
    unsigned = export_pdf(report)
    document_hash = sha256_bytes(unsigned)
    signed_at = timezone.localtime()
    signed_pdf = export_pdf(report, signature={
        "name": user.full_name or user.username,
        "signed_at": f"{signed_at:%d.%m.%Y %H:%M:%S} (UTC{signed_at:%z})",
        "hash": document_hash,
    })
    signed = SignedReport(
        commission=commission, period_start=date_from, period_end=date_to,
        document_hash=document_hash, signed_by=user, ip_address=ip,
    )
    signed.pdf.save(f"report_{commission.short_name}_{date_from:%Y%m%d}_{date_to:%Y%m%d}.pdf", ContentFile(signed_pdf), save=False)
    signed.save()
    audit.log("sign_report", signed, {"hash": document_hash, "period": [str(date_from), str(date_to)]}, user=user, ip=ip)
    return signed
