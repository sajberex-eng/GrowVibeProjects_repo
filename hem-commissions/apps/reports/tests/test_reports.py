import io
import re
import shutil
import tempfile
from datetime import date, datetime, timedelta

from django.core import mail
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from apps.committees.models import CompositionChange, LegalAct, MemberRecord, Order, Rubric
from apps.committees.testing import make_commission, make_meeting, make_user
from apps.meetings.models import AgendaItem, Attendance, Meeting, MeetingReschedule
from apps.protocols.models import Assignment, Decision, DissentingOpinion, Protocol, ProtocolItem
from apps.reports import services
from apps.reports.models import SignedReport
from apps.voting.models import VotingSession

TODAY = date(2026, 9, 16)


def aware(*args):
    return timezone.make_aware(datetime(*args))


class ReportFixture(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.commission, people = make_commission(members=3, start=date(2025, 1, 1))
        cls.other, other_people = make_commission(members=1, start=date(2025, 1, 1))
        cls.chairman = people["chairman"]
        cls.secretary = people["secretary"]
        cls.members = people["members"]
        cls.outsider = make_user()
        c = cls.commission
        m1 = make_meeting(c, starts_at=aware(2026, 2, 10, 10, 0), status=Meeting.Status.HELD)
        m2 = make_meeting(c, starts_at=aware(2026, 3, 10, 10, 0), status=Meeting.Status.HELD, kind=Meeting.Kind.EXTRA)
        make_meeting(c, starts_at=aware(2026, 4, 1, 10, 0), status=Meeting.Status.CANCELLED)
        make_meeting(c, starts_at=aware(2026, 5, 1, 10, 0))
        m_out = make_meeting(c, starts_at=aware(2026, 8, 1, 10, 0), status=Meeting.Status.HELD)
        make_meeting(cls.other, starts_at=aware(2026, 2, 11, 10, 0), status=Meeting.Status.HELD)
        MeetingReschedule.objects.create(meeting=m2, old_starts_at=aware(2026, 3, 5, 10), new_starts_at=m2.starts_at)
        S = Attendance.Status
        for user, status in [
            (cls.chairman, S.PRESENT), (cls.secretary, S.ONLINE), (cls.members[0], S.PRESENT), (cls.members[1], S.ABSENT_EXCUSED),
        ]:
            Attendance.objects.create(meeting=m1, user=user, status=status)
        Attendance.objects.create(meeting=m2, user=cls.chairman, status=S.PRESENT)
        Attendance.objects.create(meeting=m2, user=cls.members[2], status=S.ABSENT)

        rubric = Rubric.objects.create(commission=c, name="Трансфузиология")
        AgendaItem.objects.create(meeting=m1, title="В1", rubric=rubric)
        AgendaItem.objects.create(meeting=m1, title="В2", rubric=rubric, position=2)
        AgendaItem.objects.create(meeting=m2, title="В3")
        AgendaItem.objects.create(meeting=m_out, title="Вне периода", rubric=rubric)

        p1 = Protocol.objects.create(
            commission=c, meeting=m1, protocol_date=date(2026, 2, 10), status=Protocol.Status.SIGNED,
            signed_at=aware(2026, 2, 15, 12, 0), number="1",
        )
        Protocol.objects.create(
            commission=c, meeting=m2, protocol_date=date(2026, 3, 10), status=Protocol.Status.SIGNED,
            signed_at=aware(2026, 3, 20, 12, 0), number="2",
        )
        item = ProtocolItem.objects.create(protocol=p1, title="В1")
        d1 = Decision.objects.create(protocol=p1, item=item, number="1.1", text="Утвердить СОП")
        Decision.objects.create(protocol=p1, item=item, number="1.2", text="Направить на обучение")
        DissentingOpinion.objects.create(protocol=p1, item=item, author=cls.members[0], text="Не согласен")
        A = Assignment.Status
        m0, m1u = cls.members[0], cls.members[1]
        Assignment.objects.create(decision=d1, text="a1", responsible=m0, due_date=date(2026, 3, 1), status=A.DONE, completed_on=date(2026, 2, 25))
        Assignment.objects.create(decision=d1, text="a2", responsible=m0, due_date=date(2026, 3, 1), status=A.DONE_LATE, completed_on=date(2026, 3, 5))
        Assignment.objects.create(decision=d1, text="a3", responsible=m1u, due_date=date(2026, 4, 1))
        Assignment.objects.create(decision=d1, text="a4", responsible=m1u, due_date=date(2027, 1, 1))
        Assignment.objects.create(decision=d1, text="a5", responsible_name="Внешний", due_date=date(2026, 4, 1), status=A.CANCELLED)

        VotingSession.objects.create(
            commission=c, question="Заочно", deadline=aware(2026, 4, 10, 18), status=VotingSession.Status.CLOSED,
            closed_at=aware(2026, 4, 10, 18),
        )
        VotingSession.objects.create(
            commission=c, question="Заочно позже", deadline=aware(2026, 8, 10, 18), status=VotingSession.Status.CLOSED,
            closed_at=aware(2026, 8, 10, 18),
        )
        LegalAct.objects.create(commission=c, title="НПА1", status=LegalAct.Status.ACTUAL, last_checked_at=aware(2026, 9, 1, 3))
        LegalAct.objects.create(commission=c, title="НПА2", status=LegalAct.Status.ACTUAL)
        LegalAct.objects.create(commission=c, title="НПА3", status=LegalAct.Status.LOST_FORCE)
        LegalAct.objects.create(commission=c, title="НПА4", status=LegalAct.Status.NEEDS_CHECK)

        newcomer = make_user()
        order = Order.objects.create(commission=c, number="15-од", signed_on=date(2026, 5, 20), status=Order.Status.SIGNED)
        CompositionChange.objects.create(
            commission=c, action=CompositionChange.Action.ADD, user=newcomer, new_role=MemberRecord.Role.MEMBER,
            order=order, status=CompositionChange.Status.APPLIED, applied_on=date(2026, 5, 20),
        )
        MemberRecord.objects.create(commission=c, user=newcomer, role=MemberRecord.Role.MEMBER, start_date=date(2026, 5, 20), start_order=order)

    def report(self):
        return services.build_report([self.commission], date(2026, 1, 1), date(2026, 6, 30), today=TODAY)


class BuildReportTests(ReportFixture):
    def test_meeting_counts(self):
        s = self.report()["commissions"][0]
        self.assertEqual(s["meetings"], {
            "total": 4, "held": 2, "postponed": 1, "cancelled": 1, "upcoming": 1,
            "extra": 1, "extra_held": 1, "quorum_reached": 1, "quorum_not_reached": 1,
        })
        self.assertEqual(s["absentee_votings"], 1)

    def test_attendance(self):
        s = self.report()["commissions"][0]
        rows = {r["user"].pk: r for r in s["attendance"]}
        self.assertEqual(len(rows), 5)  # новичок вошёл в состав после заседаний
        chairman = rows[self.chairman.pk]
        self.assertEqual((chairman["meetings"], chairman["present"], chairman["percent"]), (2, 2, 100.0))
        secretary = rows[self.secretary.pk]
        self.assertEqual((secretary["online"], secretary["unmarked"], secretary["percent"]), (1, 1, 50.0))
        self.assertEqual(rows[self.members[1].pk]["absent_excused"], 1)
        self.assertEqual(rows[self.members[2].pk]["absent"], 1)
        self.assertEqual(rows[self.members[2].pk]["percent"], 0.0)
        self.assertEqual(s["avg_attendance"], 40.0)  # 4 из 10

    def test_protocols_and_dissents(self):
        s = self.report()["commissions"][0]
        self.assertEqual(s["protocol_timeliness"], {"signed": 2, "unsigned": 0, "avg_days": 7.5, "max_days": 10})
        self.assertEqual(s["dissents"], 1)

    def test_substantive(self):
        s = self.report()["commissions"][0]
        self.assertEqual(s["rubrics"], [{"name": "Трансфузиология", "count": 2}, {"name": "Без рубрики", "count": 1}])
        self.assertEqual(s["questions_total"], 3)
        self.assertEqual(s["decisions_count"], 2)
        self.assertEqual(s["assignments"], {
            "total": 5, "done_on_time": 1, "done_late": 1, "overdue": 1, "open": 1, "cancelled": 1,
        })
        by_name = {r["name"]: r for r in s["assignments_by_responsible"]}
        self.assertEqual(by_name[self.members[1].short_name]["overdue"], 1)
        self.assertEqual(by_name["Внешний"]["cancelled"], 1)

    def test_legal_and_composition(self):
        s = self.report()["commissions"][0]
        legal = s["legal"]
        self.assertEqual((legal["total"], legal["actual"], legal["lost_force"], legal["needs_check"]), (4, 2, 1, 1))
        self.assertEqual(legal["last_check"], aware(2026, 9, 1, 3))
        self.assertEqual(len(s["composition_changes"]), 1)
        self.assertEqual([o.number for o in s["orders"]], ["15-од"])
        self.assertEqual(s["chairman"], self.chairman)

    def test_totals_over_several_commissions(self):
        report = services.build_report([self.commission, self.other], date(2026, 1, 1), date(2026, 6, 30), today=TODAY)
        self.assertEqual(report["totals"]["meetings"]["held"], 3)
        self.assertEqual(report["totals"]["decisions_count"], 2)

    def test_period_bounds(self):
        self.assertEqual(services.period_bounds("h1", 2026), (date(2026, 1, 1), date(2026, 6, 30)))
        self.assertEqual(services.period_bounds("h2", 2026), (date(2026, 7, 1), date(2026, 12, 31)))
        self.assertEqual(services.period_bounds("year", 2026), (date(2026, 1, 1), date(2026, 12, 31)))


class ReportViewsTests(ReportFixture):
    def setUp(self):
        self.media = tempfile.mkdtemp()
        self.override = override_settings(MEDIA_ROOT=self.media)
        self.override.enable()

    def tearDown(self):
        self.override.disable()
        shutil.rmtree(self.media, ignore_errors=True)

    def params(self, **extra):
        return {"commission": self.commission.pk, "period": "h1", "year": 2026, **extra}

    def test_index_renders(self):
        self.client.force_login(self.members[0])
        response = self.client.get(reverse("reports:index"), self.params())
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Трансфузиология")
        self.assertNotContains(response, "Подписать отчёт")
        response = self.client.get(reverse("reports:index"))
        self.assertEqual(response.status_code, 200)

    def test_index_custom_period_validation(self):
        self.client.force_login(self.members[0])
        response = self.client.get(reverse("reports:index"), self.params(period="custom"))
        self.assertIsNone(response.context["report"])
        response = self.client.get(reverse("reports:index"), self.params(period="custom", date_from="2026-02-01", date_to="2026-02-28"))
        self.assertEqual(response.context["report"]["commissions"][0]["meetings"]["held"], 1)

    def test_access_limited_to_visible(self):
        self.client.force_login(self.outsider)
        response = self.client.get(reverse("reports:index"), self.params())
        self.assertIsNone(response.context["report"])  # чужая комиссия не проходит валидацию
        response = self.client.get(reverse("reports:index"), {"period": "h1", "year": 2026})
        self.assertEqual(response.context["report"]["commissions"], [])

    def test_export_xlsx(self):
        from openpyxl import load_workbook

        self.client.force_login(self.members[0])
        response = self.client.get(reverse("reports:export_xlsx"), self.params())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ["Заседания", "Посещаемость", "Вопросы", "Решения", "Поручения", "НПА", "Состав"])

    def test_export_pdf(self):
        self.client.force_login(self.members[0])
        response = self.client.get(reverse("reports:export_pdf"), {"period": "h1", "year": 2026})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_sign_flow(self):
        self.client.force_login(self.chairman)
        index = self.client.get(reverse("reports:index"), self.params())
        self.assertContains(index, "Подписать отчёт")
        mail.outbox.clear()
        response = self.client.post(reverse("reports:sign"), self.params(step="send"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(mail.outbox), 1)
        code = re.search(r"(\d{6})", mail.outbox[0].subject).group(1)
        wrong = "000000" if code != "000000" else "111111"
        response = self.client.post(reverse("reports:sign"), self.params(step="confirm", code=wrong))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(SignedReport.objects.exists())
        # после неверной попытки код остаётся действительным
        response = self.client.post(reverse("reports:sign"), self.params(step="confirm", code=code))
        self.assertEqual(response.status_code, 302)
        signed = SignedReport.objects.get()
        self.assertEqual(signed.signed_by, self.chairman)
        self.assertEqual((signed.period_start, signed.period_end), (date(2026, 1, 1), date(2026, 6, 30)))
        self.assertEqual(len(signed.document_hash), 64)

        self.client.force_login(self.members[0])
        listing = self.client.get(reverse("reports:signed_list"))
        self.assertContains(listing, signed.document_hash)
        response = self.client.get(reverse("reports:signed_download", args=[signed.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(b"".join(response.streaming_content).startswith(b"%PDF"))
        response.close()

        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(reverse("reports:signed_download", args=[signed.pk])).status_code, 403)
        self.assertNotContains(self.client.get(reverse("reports:signed_list")), signed.document_hash)

    def test_only_chairman_can_sign(self):
        self.client.force_login(self.secretary)
        response = self.client.post(reverse("reports:sign"), self.params(step="send"))
        self.assertEqual(response.status_code, 403)

    def test_sign_requires_single_commission(self):
        self.client.force_login(self.chairman)
        response = self.client.post(reverse("reports:sign"), {"period": "h1", "year": 2026, "step": "send"})
        self.assertEqual(response.status_code, 302)

    def test_sign_expired_code(self):
        from apps.accounts.models import OneTimeCode

        self.client.force_login(self.chairman)
        self.client.post(reverse("reports:sign"), self.params(step="send"))
        code = re.search(r"(\d{6})", mail.outbox[-1].subject).group(1)
        OneTimeCode.objects.update(expires_at=timezone.now() - timedelta(minutes=1))
        response = self.client.post(reverse("reports:sign"), self.params(step="confirm", code=code))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(SignedReport.objects.exists())
