"""Формирование DOCX-проекта приказа о составе комиссии.

Если администратор загрузил шаблон (DocumentTemplate вида «order»), документ
строится по нему через docxtpl. Иначе — стандартный документ через python-docx.

Переменные шаблона (docxtpl / Jinja2):
  organization, commission (name, short_name), order (title), number, date,
  preamble, add_list, remove_list, change_list, members.
  Элементы add_list / remove_list / members: num, fio, role, position, city, basis.
  Элементы change_list: num, fio, description, basis.
Пример шаблона: manage.py write_order_template → docs/templates/order_template.docx
"""
import io
from pathlib import Path

from django.conf import settings
from django.core.files.base import ContentFile
from django.utils import timezone

from .models import CompositionChange, DocumentTemplate, MemberRecord

MONTHS_GEN = [
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
]
ROLE_LABELS = dict(MemberRecord.Role.choices)
DEFAULT_PREAMBLE = (
    "В целях обеспечения деятельности комиссии и в связи с кадровыми изменениями ПРИКАЗЫВАЮ:"
)


def ru_date(value):
    if not value:
        return "«___» ____________ 20__ г."
    return f"{value.day} {MONTHS_GEN[value.month - 1]} {value.year} г."


def _role_order(role):
    return MemberRecord.ROLE_ORDER.get(role, 2)


def _person(user, role, position, city, basis=""):
    return {
        "fio": user.full_name or user.username,
        "short_fio": user.short_name,
        "role": ROLE_LABELS.get(role, role),
        "role_code": role,
        "position": position or "",
        "city": city or "",
        "basis": basis or "",
        "_last": user.last_name,
    }


def projected_composition(commission, changes, on_date=None):
    """Состав комиссии после применения изменений (без записи в БД)."""
    on_date = on_date or timezone.localdate()
    people = {}
    for rec in commission.members_on(on_date):
        people[rec.user_id] = {"user": rec.user, "role": rec.role, "position": rec.position_text, "city": rec.city_text}
    A = CompositionChange.Action
    for ch in changes:
        u = ch.user
        if ch.action == A.ADD:
            people[u.pk] = {
                "user": u,
                "role": ch.new_role or MemberRecord.Role.MEMBER,
                "position": ch.new_position or (u.position.name if u.position_id else ""),
                "city": u.city.name if u.city_id else "",
            }
        elif ch.action == A.REMOVE:
            people.pop(u.pk, None)
        elif u.pk in people:
            if ch.action == A.CHANGE_ROLE and ch.new_role:
                people[u.pk]["role"] = ch.new_role
            if ch.new_position:
                people[u.pk]["position"] = ch.new_position
    rows = [_person(p["user"], p["role"], p["position"], p["city"]) for p in people.values()]
    rows.sort(key=lambda r: (_role_order(r["role_code"]), r["_last"]))
    for i, r in enumerate(rows, 1):
        r["num"] = i
    return rows


def order_context(order):
    commission = order.commission
    changes = list(order.changes.select_related("user", "user__position", "user__city").order_by("created_at"))
    on_date = order.signed_on or timezone.localdate()
    current = {r.user_id: r for r in commission.members_on(on_date)}
    A = CompositionChange.Action
    add_list, remove_list, change_list = [], [], []
    for ch in changes:
        u = ch.user
        rec = current.get(u.pk)
        if ch.action == A.ADD:
            add_list.append(_person(
                u, ch.new_role or MemberRecord.Role.MEMBER,
                ch.new_position or (u.position.name if u.position_id else ""),
                u.city.name if u.city_id else "", ch.basis,
            ))
        elif ch.action == A.REMOVE:
            remove_list.append(_person(
                u, rec.role if rec else "", rec.position_text if rec else "", rec.city_text if rec else "", ch.basis,
            ))
        else:
            parts = []
            if ch.action == A.CHANGE_ROLE and ch.new_role:
                parts.append(f"считать {ROLE_LABELS.get(ch.new_role, ch.new_role).lower()}")
            if ch.new_position:
                parts.append(f"должность — {ch.new_position}")
            change_list.append({
                "fio": u.full_name or u.username,
                "description": "; ".join(parts),
                "basis": ch.basis,
            })
    for lst in (add_list, remove_list, change_list):
        for i, item in enumerate(lst, 1):
            item["num"] = i
    return {
        "organization": settings.ORGANIZATION_NAME,
        "commission": {"name": commission.name, "short_name": commission.short_name},
        "order": {"title": order.title},
        "title": order.title,
        "number": order.number or "____",
        "date": ru_date(order.signed_on),
        "preamble": order.preamble or DEFAULT_PREAMBLE,
        "add_list": add_list,
        "remove_list": remove_list,
        "change_list": change_list,
        "members": projected_composition(commission, changes, on_date),
    }


# ---------- рендеринг ----------

def render_from_template(template_file, context):
    from docxtpl import DocxTemplate

    with template_file.open("rb") as fh:
        data = io.BytesIO(fh.read())
    tpl = DocxTemplate(data)
    tpl.render(context)
    out = io.BytesIO()
    tpl.save(out)
    return out.getvalue()


def _base_document():
    from docx import Document
    from docx.shared import Cm, Pt

    doc = Document()
    for section in doc.sections:
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if fonts is not None:
        fonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Times New Roman")
    return doc


def _para(doc, text, bold=False, align=None, size=None, space_after=6):
    from docx.shared import Pt

    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _person_line(item):
    parts = [item["fio"]]
    if item.get("position"):
        parts.append(item["position"])
    if item.get("city"):
        parts.append(f"г. {item['city']}")
    line = " — ".join(parts)
    if item.get("role"):
        line += f" ({item['role'].lower()})"
    return line


def render_default(context):
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = _base_document()
    center = WD_ALIGN_PARAGRAPH.CENTER
    _para(doc, context["organization"], bold=True, align=center)
    _para(doc, "ПРИКАЗ", bold=True, align=center, size=16, space_after=12)
    head = doc.add_table(rows=1, cols=2)
    head.cell(0, 0).text = context["date"]
    head.cell(0, 1).text = f"№ {context['number']}"
    head.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _para(doc, "")
    _para(doc, f"{context['title']} «{context['commission']['name']}»", bold=True, space_after=12)
    _para(doc, context["preamble"], align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    point = 0
    if context["remove_list"]:
        point += 1
        _para(doc, f"{point}. Вывести из состава комиссии:")
        for item in context["remove_list"]:
            line = f"{point}.{item['num']}. {_person_line(item)}"
            if item["basis"]:
                line += f"; основание: {item['basis']}"
            _para(doc, line + ";")
    if context["add_list"]:
        point += 1
        _para(doc, f"{point}. Ввести в состав комиссии:")
        for item in context["add_list"]:
            line = f"{point}.{item['num']}. {_person_line(item)}"
            if item["basis"]:
                line += f"; основание: {item['basis']}"
            _para(doc, line + ";")
    if context["change_list"]:
        point += 1
        _para(doc, f"{point}. Внести изменения в сведения о членах комиссии:")
        for item in context["change_list"]:
            line = f"{point}.{item['num']}. {item['fio']}: {item['description']}"
            if item["basis"]:
                line += f"; основание: {item['basis']}"
            _para(doc, line + ";")
    point += 1
    _para(doc, f"{point}. Утвердить состав комиссии в редакции согласно приложению к настоящему приказу.")
    point += 1
    _para(doc, f"{point}. Контроль за исполнением настоящего приказа оставляю за собой.")
    _para(doc, "")
    _para(doc, "Руководитель  _______________________", space_after=0)

    doc.add_page_break()
    _para(doc, f"Приложение к приказу от {context['date']} № {context['number']}", align=WD_ALIGN_PARAGRAPH.RIGHT, size=11)
    _para(doc, f"Состав комиссии «{context['commission']['name']}»", bold=True, align=center, space_after=10)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ["№", "ФИО", "Роль в комиссии", "Должность", "Город"]):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(11)
    for m in context["members"]:
        cells = table.add_row().cells
        for cell, text in zip(cells, [str(m["num"]), m["fio"], m["role"], m["position"], m["city"]]):
            cell.text = ""
            cell.paragraphs[0].add_run(text).font.size = Pt(11)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def render_order_docx(order):
    context = order_context(order)
    template = DocumentTemplate.find(DocumentTemplate.Kind.ORDER, order.commission)
    if template and template.file:
        return render_from_template(template.file, context)
    return render_default(context)


def generate_order_draft(order):
    """Формирует DOCX и сохраняет его в order.draft_file."""
    data = render_order_docx(order)
    if order.draft_file:
        order.draft_file.delete(save=False)
    order.draft_file.save(f"order-{order.pk}.docx", ContentFile(data), save=True)
    return order


# ---------- образец шаблона ----------

def build_sample_order_template():
    """DOCX-шаблон с плейсхолдерами docxtpl — основа для настройки администратором."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = _base_document()
    center = WD_ALIGN_PARAGRAPH.CENTER
    _para(doc, "{{ organization }}", bold=True, align=center)
    _para(doc, "ПРИКАЗ", bold=True, align=center, size=16)
    _para(doc, "{{ date }}                                             № {{ number }}")
    _para(doc, "{{ title }} «{{ commission.name }}»", bold=True)
    _para(doc, "{{ preamble }}")
    _para(doc, "{%p if remove_list %}")
    _para(doc, "1. Вывести из состава комиссии:")
    _para(doc, "{%p for m in remove_list %}")
    _para(doc, "1.{{ m.num }}. {{ m.fio }} — {{ m.position }}{% if m.basis %}; основание: {{ m.basis }}{% endif %};")
    _para(doc, "{%p endfor %}")
    _para(doc, "{%p endif %}")
    _para(doc, "{%p if add_list %}")
    _para(doc, "2. Ввести в состав комиссии:")
    _para(doc, "{%p for m in add_list %}")
    _para(doc, "2.{{ m.num }}. {{ m.fio }} — {{ m.position }}, г. {{ m.city }} ({{ m.role|lower }});")
    _para(doc, "{%p endfor %}")
    _para(doc, "{%p endif %}")
    _para(doc, "{%p if change_list %}")
    _para(doc, "3. Внести изменения в сведения о членах комиссии:")
    _para(doc, "{%p for m in change_list %}")
    _para(doc, "3.{{ m.num }}. {{ m.fio }}: {{ m.description }};")
    _para(doc, "{%p endfor %}")
    _para(doc, "{%p endif %}")
    _para(doc, "4. Утвердить состав комиссии в редакции согласно приложению.")
    _para(doc, "5. Контроль за исполнением настоящего приказа оставляю за собой.")
    _para(doc, "Руководитель  _______________________")
    doc.add_page_break()
    _para(doc, "Приложение к приказу от {{ date }} № {{ number }}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    _para(doc, "Состав комиссии «{{ commission.name }}»", bold=True, align=center)
    table = doc.add_table(rows=4, cols=5)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["№", "ФИО", "Роль", "Должность", "Город"]):
        cell.text = text
    table.rows[1].cells[0].text = "{%tr for m in members %}"
    for cell, text in zip(table.rows[2].cells, ["{{ m.num }}", "{{ m.fio }}", "{{ m.role }}", "{{ m.position }}", "{{ m.city }}"]):
        cell.text = text
    table.rows[3].cells[0].text = "{%tr endfor %}"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def write_sample_order_template(path=None):
    path = Path(path) if path else Path(settings.BASE_DIR) / "docs" / "templates" / "order_template.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(build_sample_order_template())
    return path
