import io
from pathlib import Path

from django.core.files.base import ContentFile
from docx import Document

from apps.committees import docgen
from apps.committees.models import CompositionChange, DocumentTemplate, MemberRecord, Order
from apps.committees.testing import make_user

from .base import CommissionTestCase


def docx_text(data):
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


class DocgenTests(CommissionTestCase):
    def setUp(self):
        super().setUp()
        self.newbie = make_user(last_name="Новикова", first_name="Алия")
        self.order = Order.objects.create(commission=self.commission, status=Order.Status.DRAFT, title="О составе комиссии")
        CompositionChange.objects.create(
            commission=self.commission, action=CompositionChange.Action.ADD, user=self.newbie,
            new_role=MemberRecord.Role.MEMBER, basis="заявление", order=self.order,
        )
        CompositionChange.objects.create(
            commission=self.commission, action=CompositionChange.Action.REMOVE, user=self.members[1],
            order=self.order,
        )

    def test_context(self):
        ctx = docgen.order_context(self.order)
        self.assertEqual([m["fio"] for m in ctx["add_list"]], [self.newbie.full_name])
        self.assertEqual([m["fio"] for m in ctx["remove_list"]], [self.members[1].full_name])
        fios = [m["fio"] for m in ctx["members"]]
        self.assertIn(self.newbie.full_name, fios)
        self.assertNotIn(self.members[1].full_name, fios)
        self.assertEqual(ctx["members"][0]["fio"], self.chairman.full_name)
        self.assertEqual(len(fios), 5)  # 2 + 3 члена + 1 новый − 1 выведенный

    def test_default_docx(self):
        docgen.generate_order_draft(self.order)
        self.order.refresh_from_db()
        self.assertTrue(self.order.draft_file)
        with self.order.draft_file.open("rb") as fh:
            text = docx_text(fh.read())
        self.assertIn("ПРИКАЗ", text)
        self.assertIn("Вывести из состава", text)
        self.assertIn("Ввести в состав", text)
        self.assertIn("Новикова Алия", text)
        self.assertIn(self.chairman.full_name, text)

    def test_custom_template(self):
        tpl_bytes = docgen.build_sample_order_template()
        DocumentTemplate.objects.create(
            kind=DocumentTemplate.Kind.ORDER, file=ContentFile(tpl_bytes, name="tpl.docx"),
        )
        text = docx_text(docgen.render_order_docx(self.order))
        self.assertIn("Новикова Алия", text)
        self.assertIn(self.chairman.full_name, text)
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)

    def test_write_sample_template(self):
        path = Path(self._media) / "sample" / "order_template.docx"
        docgen.write_sample_order_template(path)
        self.assertIn("{{ organization }}", docx_text(path.read_bytes()))
