"""Экспорт протокола в DOCX: по шаблону DocumentTemplate (docxtpl) или по умолчанию (python-docx).

Контекст шаблона: всё из снимка (organization, commission.name, number, date,
place, format, preamble, attendance.present/absent/invited, quorum, items[...],
dissents[...], transcription_used, officers) плюс date_text, title, transcription_remark,
assignments (плоский список), signatures, frozen_hash, signed_hash, verify_url.
"""
import io
from datetime import date

from django.utils import timezone

from apps.committees.models import DocumentTemplate

from . import pdf
from .models import Protocol


def build_context(protocol, snap):
    ctx = dict(snap)
    ctx["date_text"] = pdf._date(snap.get("date"))
    kind_title = "заочного голосования" if snap.get("kind") == Protocol.Kind.ABSENTEE else "заседания"
    ctx["title"] = f"ПРОТОКОЛ № {snap.get('number') or 'б/н'} {kind_title} комиссии «{snap['commission']['name']}»"
    ctx["transcription_remark"] = pdf.TRANSCRIPTION_REMARK if snap.get("transcription_used") else ""
    flat = []
    for item in snap.get("items", []):
        for d in item.get("decisions", []):
            for a in d.get("assignments", []):
                flat.append({**a, "due_date_text": pdf._date(a.get("due_date")), "item_position": item["position"]})
    ctx["assignments"] = flat
    ctx["signatures"] = [
        {"full_name": s.full_name, "role": s.get_role_display(), "position": s.position_text,
         "signed_at": timezone.localtime(s.signed_at), "method": pdf.SIGN_METHOD}
        for s in protocol.active_signatures()
    ]
    ctx["frozen_hash"] = protocol.frozen_hash
    ctx["signed_hash"] = protocol.signed_hash
    ctx["verify_url"] = protocol.verify_url
    ctx["status"] = protocol.get_status_display()
    return ctx


def _render_template(template, ctx):
    from docxtpl import DocxTemplate

    with template.file.open("rb") as fh:
        source = io.BytesIO(fh.read())
    doc = DocxTemplate(source)
    doc.render(ctx)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _render_default(ctx):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    def center(text, bold=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = bold
        return p

    def label(text):
        p = doc.add_paragraph()
        p.add_run(text).bold = True

    center(ctx.get("organization", ""))
    center(ctx["title"], bold=True)
    doc.add_paragraph(f"Дата: {ctx['date_text']}")
    if ctx.get("time"):
        doc.add_paragraph(f"Время начала: {ctx['time']}")
    if ctx.get("place"):
        doc.add_paragraph(f"Место проведения: {ctx['place']}")
    if ctx.get("format"):
        doc.add_paragraph(f"Форма проведения: {ctx['format']}")
    if ctx.get("preamble"):
        doc.add_paragraph(ctx["preamble"])

    att = ctx.get("attendance") or {}
    absentee = ctx.get("kind") == Protocol.Kind.ABSENTEE
    for key, title in [
        ("present", "Приняли участие в голосовании:" if absentee else "Присутствовали:"),
        ("absent", "Не приняли участие в голосовании:" if absentee else "Отсутствовали:"),
        ("invited", "Приглашённые:"),
    ]:
        people = att.get(key) or []
        if people:
            label(title)
            doc.add_paragraph(pdf._people_lines(people))
    q = ctx.get("quorum")
    if q:
        state = "Кворум имеется" if q.get("reached") else "Кворум отсутствует"
        doc.add_paragraph(f"{state}: присутствует {q['present']} из {q['total']} (требуется {q['needed']}).")

    label("ПОВЕСТКА ДНЯ:")
    for item in ctx.get("items", []):
        extra = f" (ID пациента: {item['patient_id']})" if item.get("patient_id") else ""
        doc.add_paragraph(f"{item['position']}. {item['title']}{extra}")

    for item in ctx.get("items", []):
        doc.add_heading(f"{item['position']}. {item['title']}", level=2)
        if item.get("heard"):
            label("СЛУШАЛИ:")
            doc.add_paragraph(item["heard"])
        if item.get("discussed"):
            label("ВЫСТУПИЛИ:")
            doc.add_paragraph(item["discussed"])
        if item.get("resolved") or item.get("decisions") or item.get("dissent_note"):
            label("РЕШИЛИ:")
            if item.get("resolved"):
                doc.add_paragraph(item["resolved"])
            for d in item.get("decisions", []):
                doc.add_paragraph((f"{d['number']}. " if d.get("number") else "— ") + d["text"])
            if item.get("dissent_note"):
                doc.add_paragraph(item["dissent_note"])
        if item.get("vote_summary"):
            label("Результаты голосования:")
            doc.add_paragraph(item["vote_summary"])

    if ctx.get("dissents"):
        doc.add_heading("Особые мнения членов комиссии", level=2)
        for d in ctx["dissents"]:
            label(f"По вопросу {d['item_position']} — {d['author']}")
            if d.get("text"):
                doc.add_paragraph(d["text"])
            if d.get("file_name"):
                doc.add_paragraph(f"Приложение: {d['file_name']}")

    if ctx["assignments"]:
        doc.add_heading("Поручения", level=2)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ["№", "Поручение", "Ответственный", "Срок"]):
            cell.text = text
        for n, a in enumerate(ctx["assignments"], 1):
            cells = table.add_row().cells
            cells[0].text = str(n)
            cells[1].text = a["text"]
            cells[2].text = a["responsible"] + (
                "\nСоисполнители: " + ", ".join(a["co_executors"]) if a.get("co_executors") else ""
            )
            cells[3].text = a["due_date_text"]

    if ctx["transcription_remark"]:
        doc.add_paragraph(ctx["transcription_remark"])
    officers = ctx.get("officers") or {}
    doc.add_paragraph("")
    doc.add_paragraph(f"Председатель комиссии  _______________  {officers.get('chairman', '')}")
    doc.add_paragraph(f"Секретарь комиссии  _______________  {officers.get('secretary', '')}")
    if ctx["signatures"]:
        label("Подписано простой электронной подписью:")
        for s in ctx["signatures"]:
            doc.add_paragraph(f"{s['role']}: {s['full_name']} — {s['signed_at']:%d.%m.%Y %H:%M}")
        doc.add_paragraph(f"SHA-256: {ctx['frozen_hash']}")
        doc.add_paragraph(f"Проверка: {ctx['verify_url']}")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def export_docx(protocol, snap):
    ctx = build_context(protocol, snap)
    kind = (
        DocumentTemplate.Kind.ABSENTEE_PROTOCOL if protocol.kind == Protocol.Kind.ABSENTEE
        else DocumentTemplate.Kind.PROTOCOL
    )
    template = DocumentTemplate.find(kind, protocol.commission)
    if template is None and kind == DocumentTemplate.Kind.ABSENTEE_PROTOCOL:
        template = DocumentTemplate.find(DocumentTemplate.Kind.PROTOCOL, protocol.commission)
    if template is not None:
        return _render_template(template, ctx)
    return _render_default(ctx)


def filename(protocol, ext):
    number = (protocol.number or "bn").replace("/", "-").replace(" ", "_")
    return f"protocol_{number}_{protocol.protocol_date or date.today():%Y%m%d}.{ext}"
